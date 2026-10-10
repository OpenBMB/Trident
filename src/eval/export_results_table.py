#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
export_results_table.py

Export evaluation results into a spreadsheet (.xlsx) that opens in Excel / WPS.

Several metrics are supported, e.g.

    --metric ndcg@10,recall@10,mrr@10

Each metric gets its own sheet.

Usage:

    # single metric
    python src/eval/export_results_table.py \
        --eval_root eval/xxx \
        --embed_root embedding/xxx \
        --model_name Trident-Qwen3VL-2B \
        --metric ndcg@10

    # several metrics
    python src/eval/export_results_table.py \
        --eval_root eval/xxx \
        --embed_root embedding/xxx \
        --model_name Trident-Qwen3VL-2B \
        --dims 128,256,512,1024,2048 \
        --datasets ChartQA,DocVQA,InfoVQA,SlideVQA,ViDoSeek,Dude \
        --metric ndcg@10,recall@10,mrr@10
"""

import argparse
import glob
import json
import os
import re
from typing import Optional

PREFERRED_DATASET_ORDER = [
    "ChartQA", "DocVQA", "InfoVQA", "SlideVQA", "ViDoSeek", "Dude",
]

SETTINGS = ("Origin", "Mix")


def parse_int_list(text: str) -> list[int]:
    parts = re.split(r"[,\s\[\]]+", text.strip())
    return sorted({int(p) for p in parts if p})


def parse_str_list(text: str) -> list[str]:
    """Parse a comma-separated string, preserving order and removing duplicates."""
    return list(
        dict.fromkeys(
            part.strip()
            for part in text.split(",")
            if part.strip()
        )
    )


def load_json(path: str):
    if not os.path.isfile(path):
        return None
    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


# ---------------------------------------------------------------------------
# reading results
# ---------------------------------------------------------------------------


def read_result_dir(
    result_dir: str,
    metric: str,
) -> dict[str, dict[str, float]]:
    """
    Read one evaluation result directory and return:

        {
            "Origin": {dataset: value},
            "Mix": {dataset: value},
        }

    Values are raw fractions in [0, 1].
    """
    out: dict[str, dict[str, float]] = {
        "Origin": {},
        "Mix": {},
    }

    if not os.path.isdir(result_dir):
        return out

    # Origin
    for name in sorted(os.listdir(result_dir)):
        if not name.endswith("_result.json") or name == "mix_result.json":
            continue

        payload = load_json(os.path.join(result_dir, name))

        if not payload or "metrics" not in payload:
            continue

        dataset = payload.get("dataset") or name[:-len("_result.json")]

        out["Origin"][dataset] = fetch_metric(
            payload["metrics"],
            metric,
            name,
        )

    # Mix
    mix_payload = load_json(
        os.path.join(result_dir, "mix_result.json")
    )

    if mix_payload:
        for dataset, item in (mix_payload.get("datasets") or {}).items():
            out["Mix"][dataset] = fetch_metric(
                item["metrics"],
                metric,
                f"mix_result.json[{dataset}]",
            )

    return out


def fetch_metric(
    metrics: dict,
    metric: str,
    where: str,
) -> float:
    if metric not in metrics:
        raise SystemExit(
            f"[error] metric {metric!r} not found in {where}; "
            f"available: {sorted(metrics)}. "
            f"(The k of the metric comes from --k_values / K_VALUES at evaluation time and must be included.)"
        )

    return float(metrics[metric])


def detect_full_dim(embed_root: Optional[str]) -> Optional[int]:
    """Read the full embedding dimension from any <tag>_queries/embeddings.npy."""
    if not embed_root:
        return None

    import numpy as np

    for path in sorted(
        glob.glob(
            os.path.join(
                embed_root,
                "*_queries",
                "embeddings.npy",
            )
        )
    ):
        try:
            return int(
                np.load(
                    path,
                    mmap_mode="r",
                ).shape[1]
            )
        except Exception:  # noqa: BLE001
            continue

    return None


# ---------------------------------------------------------------------------
# table assembly
# ---------------------------------------------------------------------------


def order_datasets(
    found: set[str],
    requested: Optional[list[str]],
) -> list[str]:

    if requested:
        ordered = list(dict.fromkeys(requested))
        extra = sorted(found - set(ordered))

        if extra:
            print(
                f"[warn] the results contain datasets not listed in --datasets; "
                f"appending them at the end: {extra}"
            )

        return ordered + extra

    head = [
        d
        for d in PREFERRED_DATASET_ORDER
        if d in found
    ]

    tail = sorted(found - set(head))

    return head + tail


def to_percent(value: float) -> float:
    return round(value * 100.0, 2)


def build_rows(
    row_sources: list[
        tuple[
            str,
            dict[str, dict[str, float]],
        ]
    ],
    datasets: list[str],
) -> list[
    tuple[
        str,
        dict[
            str,
            dict[str, Optional[float]],
        ],
    ]
]:

    rows = []

    for label, per_setting in row_sources:

        row: dict[
            str,
            dict[str, Optional[float]],
        ] = {}

        for setting in SETTINGS:

            values = per_setting.get(setting, {})

            cells: dict[
                str,
                Optional[float],
            ] = {
                d: (
                    to_percent(values[d])
                    if d in values
                    else None
                )
                for d in datasets
            }

            if (
                datasets
                and all(d in values for d in datasets)
            ):
                cells["Avg"] = to_percent(
                    sum(values[d] for d in datasets)
                    / len(datasets)
                )

            else:
                cells["Avg"] = None

                if values:
                    missing = [
                        d
                        for d in datasets
                        if d not in values
                    ]

                    print(
                        f"[warn] {label} / {setting}: "
                        f"missing results for {missing}; Avg left empty"
                    )

            row[setting] = cells

        rows.append(
            (
                label,
                row,
            )
        )

    return rows


def print_table(
    rows,
    datasets: list[str],
    model_name: str,
    metric: str,
) -> None:

    columns = datasets + ["Avg"]

    header = (
        ["Model"]
        + [
            f"{c}-{s}"
            for c in columns
            for s in SETTINGS
        ]
    )

    table = [header]

    for label, row in rows:

        line = [label]

        for c in columns:
            for s in SETTINGS:

                v = row[s].get(c)

                line.append(
                    ""
                    if v is None
                    else f"{v:.2f}"
                )

        table.append(line)

    widths = [
        max(
            len(r[i])
            for r in table
        )
        for i in range(len(header))
    ]

    print("")
    print(
        f"[table] {model_name} - {metric} (%)"
    )

    for idx, r in enumerate(table):

        print(
            "  ".join(
                cell.ljust(widths[i])
                for i, cell in enumerate(r)
            )
        )

        if idx == 0:
            print(
                "  ".join(
                    "-" * w
                    for w in widths
                )
            )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def write_metric_sheet(
    ws,
    rows,
    datasets: list[str],
    model_name: str,
) -> None:

    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter

    columns = datasets + ["Avg"]

    n_cols = 1 + 2 * len(columns)

    thin = Side(
        style="thin",
        color="D0D0D0",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    center = Alignment(
        horizontal="center",
        vertical="center",
    )

    label_fill = PatternFill(
        "solid",
        fgColor="FFFFEE",
    )

    def style_range(
        r1,
        c1,
        r2,
        c2,
        **kw,
    ):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):

                cell = ws.cell(
                    row=r,
                    column=c,
                )

                cell.border = border
                cell.alignment = center

                if "font" in kw:
                    cell.font = kw["font"]

                if "fill" in kw:
                    cell.fill = kw["fill"]

    # -------------------------------------------------------
    # header
    # -------------------------------------------------------

    ws.cell(
        row=1,
        column=1,
        value="Model",
    )

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=2,
        end_column=1,
    )

    for i, name in enumerate(columns):

        c = 2 + 2 * i

        ws.cell(
            row=1,
            column=c,
            value=name,
        )

        ws.merge_cells(
            start_row=1,
            start_column=c,
            end_row=1,
            end_column=c + 1,
        )

        ws.cell(
            row=2,
            column=c,
            value="Origin",
        )

        ws.cell(
            row=2,
            column=c + 1,
            value="Mix",
        )

    style_range(
        1,
        1,
        2,
        n_cols,
        font=Font(bold=True),
    )

    # -------------------------------------------------------
    # model name
    # -------------------------------------------------------

    ws.cell(
        row=3,
        column=1,
        value=model_name,
    )

    ws.merge_cells(
        start_row=3,
        start_column=1,
        end_row=3,
        end_column=n_cols,
    )

    style_range(
        3,
        1,
        3,
        n_cols,
        font=Font(bold=True),
    )

    # -------------------------------------------------------
    # data
    # -------------------------------------------------------

    for r_off, (label, row) in enumerate(rows):

        r = 4 + r_off

        ws.cell(
            row=r,
            column=1,
            value=label,
        )

        for i, name in enumerate(columns):

            for j, setting in enumerate(SETTINGS):

                cell = ws.cell(
                    row=r,
                    column=2 + 2 * i + j,
                )

                value = row[setting].get(name)

                if value is not None:
                    cell.value = value
                    cell.number_format = "0.00"

        style_range(
            r,
            1,
            r,
            n_cols,
        )

        ws.cell(
            row=r,
            column=1,
        ).fill = label_fill

    # -------------------------------------------------------
    # sizes
    # -------------------------------------------------------

    ws.column_dimensions["A"].width = 16

    for c in range(
        2,
        n_cols + 1,
    ):
        ws.column_dimensions[
            get_column_letter(c)
        ].width = 11

    ws.freeze_panes = "B4"


def safe_sheet_name(
    metric: str,
    used_names: set[str],
) -> str:
    """
    Excel sheet names must not contain:
        : \\ / ? * [ ]

    and are at most 31 characters long.
    """

    name = re.sub(
        r'[:\\/?*\[\]]',
        "_",
        metric,
    )

    name = name[:31] or "metric"

    base = name
    idx = 2

    while name in used_names:

        suffix = f"_{idx}"

        name = (
            base[:31 - len(suffix)]
            + suffix
        )

        idx += 1

    used_names.add(name)

    return name


def write_xlsx(
    path: str,
    metric_results: dict[
        str,
        tuple[
            list,
            list[str],
        ],
    ],
    model_name: str,
) -> None:

    from openpyxl import Workbook

    wb = Workbook()

    # remove the default sheet
    default_ws = wb.active
    wb.remove(default_ws)

    used_names: set[str] = set()

    # -------------------------------------------------------
    # one sheet per metric
    # -------------------------------------------------------

    for metric, (
        rows,
        datasets,
    ) in metric_results.items():

        sheet_name = safe_sheet_name(
            metric,
            used_names,
        )

        ws = wb.create_sheet(
            sheet_name
        )

        write_metric_sheet(
            ws,
            rows,
            datasets,
            model_name,
        )

    # -------------------------------------------------------
    # notes
    # -------------------------------------------------------

    notes = wb.create_sheet("notes")

    notes.append(
        [
            "metrics",
            ", ".join(metric_results.keys()),
        ]
    )

    notes.append(
        [
            "unit",
            "percent (metric x 100), 2 decimals",
        ]
    )

    notes.append(
        [
            "Origin",
            "per-dataset evaluation, corpus = origin_corpus",
        ]
    )

    notes.append(
        [
            "Mix",
            "Mix evaluation: the mix_corpus of all datasets merged into one corpus",
        ]
    )

    notes.append(
        [
            "Avg",
            "arithmetic mean over the datasets of the row; empty if any dataset is missing",
        ]
    )

    notes.column_dimensions["A"].width = 12
    notes.column_dimensions["B"].width = 70

    os.makedirs(
        os.path.dirname(
            os.path.abspath(path)
        )
        or ".",
        exist_ok=True,
    )

    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main() -> None:

    parser = argparse.ArgumentParser(
        description="Export evaluation results into an Excel/WPS spreadsheet"
    )

    parser.add_argument(
        "--eval_root",
        required=True,
        help="eval/<RUN> directory",
    )

    parser.add_argument(
        "--embed_root",
        default=None,
        help=(
            "embedding/<RUN> directory, "
            "only used to read the model's full dimension for row labels (optional)"
        ),
    )

    parser.add_argument(
        "--model_name",
        default="model",
        help="model name shown in the table",
    )

    parser.add_argument(
        "--dims",
        default=None,
        help=(
            "Matryoshka dims, e.g. "
            "128,256,512. "
            "If omitted, a single row (full model dimension) is written"
        ),
    )

    parser.add_argument(
        "--metric",
        default="ndcg@10",
        help=(
            "evaluation metric(s). "
            "Separate several metrics with commas, e.g. "
            "ndcg@10,recall@10,mrr@10"
        ),
    )

    parser.add_argument(
        "--datasets",
        default=None,
        help=(
            "dataset column order, comma separated. "
            "Defaults to ChartQA,DocVQA,InfoVQA,"
            "SlideVQA,ViDoSeek,Dude first, "
            "then the rest alphabetically"
        ),
    )

    parser.add_argument(
        "--output",
        default=None,
        help="default: <eval_root>/results_table.xlsx",
    )

    args = parser.parse_args()

    # -------------------------------------------------------
    # parse metrics
    # -------------------------------------------------------

    metrics = parse_str_list(args.metric)

    if not metrics:
        raise SystemExit(
            "[error] --metric needs at least one metric"
        )

    requested = (
        [
            d.strip()
            for d in args.datasets.split(",")
            if d.strip()
        ]
        if args.datasets
        else None
    )

    # the full dimension only needs to be detected once
    full_dim = (
        detect_full_dim(args.embed_root)
        if not args.dims
        else None
    )

    dims = (
        parse_int_list(args.dims)
        if args.dims
        else None
    )

    # -------------------------------------------------------
    # each metric separately
    # -------------------------------------------------------

    metric_results = {}

    for metric in metrics:

        print("")
        print("=" * 80)
        print(f"[metric] {metric}")
        print("=" * 80)

        row_sources: list[
            tuple[
                str,
                dict[str, dict[str, float]],
            ]
        ] = []

        if dims:

            for dim in dims:

                result_dir = os.path.join(
                    args.eval_root,
                    "matryoshka",
                    f"dim_{dim}",
                )

                if not os.path.isdir(result_dir):

                    print(
                        f"[warn] {result_dir} does not exist "
                        f"(probably larger than the model's full dimension and skipped); "
                        f"no row for it"
                    )

                    continue

                row_sources.append(
                    (
                        f"Dim {dim}",
                        read_result_dir(
                            result_dir,
                            metric,
                        ),
                    )
                )

        else:

            label = (
                f"Dim {full_dim}"
                if full_dim
                else "Full dim"
            )

            row_sources.append(
                (
                    label,
                    read_result_dir(
                        args.eval_root,
                        metric,
                    ),
                )
            )

        # ---------------------------------------------------
        # find datasets
        # ---------------------------------------------------

        found: set[str] = set()

        for _, per_setting in row_sources:
            for values in per_setting.values():
                found.update(values)

        if not row_sources or not found:

            raise SystemExit(
                f"[error] no exportable evaluation results "
                f"for metric={metric!r} "
                f"found under {args.eval_root}"
            )

        datasets = order_datasets(
            found,
            requested,
        )

        rows = build_rows(
            row_sources,
            datasets,
        )

        print_table(
            rows,
            datasets,
            args.model_name,
            metric,
        )

        metric_results[metric] = (
            rows,
            datasets,
        )

    # -------------------------------------------------------
    # write Excel
    # -------------------------------------------------------

    output = (
        args.output
        or os.path.join(
            args.eval_root,
            "results_table.xlsx",
        )
    )

    write_xlsx(
        output,
        metric_results,
        args.model_name,
    )

    print("")
    print(
        f"[done] table written to: {output}"
    )

    print(
        f"[done] metrics: {', '.join(metrics)}"
    )


if __name__ == "__main__":
    main()